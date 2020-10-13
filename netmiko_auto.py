from netmiko import ConnectHandler
import openpyxl
def Read_excel(file_name):
    wb = openpyxl.load_workbook(file_name)
    sheet = wb.get_sheet_by_name('Sheet1')
    row = sheet.max_row
    column = sheet.max_column
    device_list = {}
    for i in range(2,row+1):
        device_list['device{0}'.format(i-1)] = []
        for j in range(1,column+1):
            vla = sheet.cell(row = i,column = j).value
            device_list['device{0}'.format(i-1)].append(vla)
    return device_list

def Connect(device,command_file):
    F = True 
    print('正在连接{0}\n'.format(device['host']))
    net_connect = ConnectHandler(**device)
    net_connect.enable()  #????

    #读取配置文件，配置文件存放在当前目录下
    for i in open(command_file,'r'):
        cmd = i.replace('\n',' ')
        print('正在执行命令：',cmd)
        result = net_connect.send_command(cmd)
        print(result)
    net_connect.disconnect() 

def Huawei(ip):  
    huawei = {
        'device_type':'huawei',
        'host':ip,
        'username':'admin',
        'password':'HUAwei@Password',
    }
    return huawei
def H3c(ip): #如果是display 华三需要刷入额外命令，改变屏幕显示长度
    h3c = {
        'device_type':'hp_comware',
        'host':ip,
        'username':'admin',
        'password':'H3c@Password',
    }
    return h3c

def Command_file(device_list,device_name):
    filename =device_list[device_name][2]
    if filename != None:
        return filename
    else:
        return 'command.txt'
def Init(ip_file):
    device_list = Read_excel(ip_file)
    for device_name in device_list.keys():
        try:
            if device_list[device_name][1] == 'huawei':
                device = Huawei(device_list[device_name][0])
                commandfile = Command_file(device_list,device_name)
                Connect(device,commandfile)
            elif device_list[device_name][1] == 'h3c': #注意华三设备在执行命令时不会一下显示所有内容。
                device = H3c(device_list[device_name][0])
                commandfile = Command_file(device_list,device_name)
                Connect(device,commandfile)
            else:
                print('未定义设备！')
                pass
        except Exception as e:
            print('连接超时：',e)
            pass 
        time.sleep(1)

if __name__ == '__main__':
    Init('ip_add.xlsx')
    
